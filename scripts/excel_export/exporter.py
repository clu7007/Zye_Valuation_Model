"""Excel workbook exporter for quarterly financial statements."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SECTION_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN_GRAY = Side(style="thin", color="D9D9D9")
THICK_BLACK = Side(style="thick", color="000000")
NEGATIVE_RED = "C00000"


def export_quarterly_financials(
    stock_id: str,
    income_statement: pd.DataFrame,
    qoq_growth: pd.DataFrame,
    yoy_growth: pd.DataFrame,
    margins: pd.DataFrame,
    output_path: Path,
) -> Path:
    """Export the full quarterly financial workbook to Excel."""

    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Quarterly Financials"
    worksheet.freeze_panes = "B4"
    worksheet.sheet_view.showGridLines = False

    max_column = income_statement.shape[1] + 1
    title_range = f"A1:{get_column_letter(max_column)}1"
    worksheet.merge_cells(title_range)
    title_cell = worksheet["A1"]
    title_cell.value = f"{stock_id} Quarterly Financials"
    title_cell.font = Font(bold=True, size=16, color="FFFFFF")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = TITLE_FILL
    worksheet.row_dimensions[1].height = 26

    worksheet["A2"] = "Unit: NT$ million; EPS: NT$ per share; sample data only"
    worksheet["A2"].font = Font(italic=True, color="666666")

    current_row = 4
    current_row = _write_section(
        worksheet=worksheet,
        start_row=current_row,
        title="Quarterly income statement",
        data=income_statement,
        value_format="#,##0",
        eps_format="#,##0.00",
        apply_negative_red=False,
    )
    current_row = _write_section(
        worksheet=worksheet,
        start_row=current_row,
        title="QoQ growth",
        data=qoq_growth,
        value_format="0.0%",
        eps_format="0.0%",
        apply_negative_red=True,
    )
    current_row = _write_section(
        worksheet=worksheet,
        start_row=current_row,
        title="YoY growth",
        data=yoy_growth,
        value_format="0.0%",
        eps_format="0.0%",
        apply_negative_red=True,
    )
    _write_section(
        worksheet=worksheet,
        start_row=current_row,
        title="Profitability margins",
        data=margins,
        value_format="0.0%",
        eps_format="0.0%",
        apply_negative_red=True,
    )

    _set_print_and_widths(worksheet, max_column)
    workbook.save(output_path)
    return output_path


def _write_section(
    worksheet,
    start_row: int,
    title: str,
    data: pd.DataFrame,
    value_format: str,
    eps_format: str,
    apply_negative_red: bool,
) -> int:
    max_column = data.shape[1] + 1

    section_cell = worksheet.cell(row=start_row, column=1, value=title)
    section_cell.font = Font(bold=True, color="000000")
    section_cell.fill = SECTION_FILL
    section_cell.alignment = Alignment(horizontal="left")
    for column in range(1, max_column + 1):
        cell = worksheet.cell(row=start_row, column=column)
        cell.fill = SECTION_FILL
        cell.border = Border(top=THICK_BLACK)

    header_row = start_row + 1
    worksheet.cell(row=header_row, column=1, value="Financial item")
    for column_index, quarter in enumerate(data.columns, start=2):
        worksheet.cell(row=header_row, column=column_index, value=quarter)

    for cell in worksheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(top=THIN_GRAY, bottom=THIN_GRAY)

    first_data_row = header_row + 1
    for row_offset, (item, values) in enumerate(data.iterrows()):
        row_number = first_data_row + row_offset
        label_cell = worksheet.cell(row=row_number, column=1, value=item)
        label_cell.alignment = Alignment(horizontal="left")
        label_cell.border = Border(bottom=THIN_GRAY)

        for column_offset, value in enumerate(values, start=2):
            cell = worksheet.cell(row=row_number, column=column_offset)
            cell.value = None if pd.isna(value) else float(value)
            cell.number_format = eps_format if item == "EPS" else value_format
            cell.alignment = Alignment(horizontal="right")
            cell.border = Border(bottom=THIN_GRAY)

    if apply_negative_red:
        data_range = (
            f"B{first_data_row}:"
            f"{get_column_letter(max_column)}{first_data_row + len(data.index) - 1}"
        )
        worksheet.conditional_formatting.add(
            data_range,
            CellIsRule(operator="lessThan", formula=["0"], font=Font(color=NEGATIVE_RED)),
        )

    return first_data_row + len(data.index) + 2


def _set_print_and_widths(worksheet, max_column: int) -> None:
    for column in range(1, max_column + 1):
        column_letter = get_column_letter(column)
        max_length = 0
        for cell in worksheet[column_letter]:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))

        if column == 1:
            width = min(max(max_length + 2, 18), 48)
        else:
            width = min(max(max_length + 2, 11), 16)
        worksheet.column_dimensions[column_letter].width = width

    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
